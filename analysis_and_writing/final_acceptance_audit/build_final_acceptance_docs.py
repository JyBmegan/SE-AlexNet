from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path("/Users/Megan/0_MyFolders/0_Projects/投稿SEAlexNet/最后修改")
MAIN_SOURCE = Path("/tmp/xb26_main_clean.docx")
RESPONSE_SOURCE = ROOT / "xb26-0201-审稿意见及说明.docx"

MAIN_OUT = ROOT / "xb26-0201_正文_作者核定版_勿直接提交.docx"
RESPONSE_OUT = ROOT / "xb26-0201-审稿意见及说明_录用后补充工作稿.docx"
REPORT_OUT = ROOT / "xb26-0201_最终核查报告_2026-10-14.docx"


def set_run_font(run, east_asia: str, size: float, *, bold: bool | None = None,
                 italic: bool | None = None, color: RGBColor | None = None) -> None:
    # Arial Unicode MS is available on the authoring Mac and gives internal
    # working documents a dependable Chinese fallback. The journal manuscript
    # continues to use the exact requested Chinese font names plus TNR.
    working_font = "Arial Unicode MS" if east_asia == "Hiragino Sans GB" else "Times New Roman"
    run.font.name = working_font
    run._element.get_or_add_rPr().rFonts.set(
        qn("w:eastAsia"), "Arial Unicode MS" if east_asia == "Hiragino Sans GB" else east_asia
    )
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def style_paragraph(paragraph, east_asia: str, size: float, *, bold: bool | None = None,
                    alignment=None) -> None:
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    if alignment is not None:
        paragraph.alignment = alignment
    for run in paragraph.runs:
        set_run_font(run, east_asia, size, bold=bold)
        run.font.highlight_color = None
        run.font.color.rgb = RGBColor(0, 0, 0)


def replace_paragraph(paragraph, text: str) -> None:
    paragraph.text = text


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def find_paragraph(doc: Document, needle: str):
    for paragraph in doc.paragraphs:
        if needle in paragraph.text:
            return paragraph
    raise ValueError(f"Paragraph not found: {needle}")


def add_comment(doc: Document, needle: str, text: str) -> None:
    paragraph = find_paragraph(doc, needle)
    if not paragraph.runs:
        paragraph.add_run(paragraph.text)
    doc.add_comment(
        runs=paragraph.runs,
        text=text,
        author="终稿核查",
        initials="QA",
    )


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_page(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)


def edit_main_manuscript() -> None:
    doc = Document(MAIN_SOURCE)
    configure_page(doc)

    replacements = {
        "关键词：表征压缩，卷积神经网络；视知觉；面孔认知；认知机制":
            "关键词 表征压缩，卷积神经网络，视知觉，面孔认知，认知机制",
        "Convolutional Neural. Networks, CNNs": "convolutional neural networks, CNNs",
        "NVIDIA GTX 3090 Ti": "NVIDIA RTX 3090 Ti",
        "0%, 20%, 30%, 40%, 50%, 60%, 70%, 80% 和 100%":
            "0%、20%、30%、40%、50%、60%、70%、80%和100%",
        "在语义表征后期（FC7FC8之间）": "在语义表征后期（FC7与FC8之间）",
        "均方误差（RMSE）": "均方根误差（RMSE）",
        "偏离真是人类知觉": "偏离真实人类知觉",
        "t(39) = 11.11. p < 0.001": "t(39) = 11.11, p < 0.001",
        "（; Mdiff = -0.08": "（Mdiff = −0.08",
        "后序压缩 （即实验1": "后序压缩（即实验1",
        "Chang & Tsao, 2017;Leopold": "Chang & Tsao, 2017; Leopold",
        "National Yang-Ming University, Taipei, Taiwan, 220.":
            "National Yang-Ming University, Taipei, Taiwan.",
    }
    for paragraph in doc.paragraphs:
        new_text = paragraph.text
        for old, new in replacements.items():
            new_text = new_text.replace(old, new)
        if new_text != paragraph.text:
            replace_paragraph(paragraph, new_text)

    replace_paragraph(doc.paragraphs[9],
        "人类高效的面孔认知可能依赖对复杂视觉信息的压缩与选择性增强。受灵长类面孔加工区低维表征证据启发，本研究以卷积神经网络（CNN）为模型系统，考察压缩激励（squeeze-and-excitation, SE）模块能否提高AlexNet输出与人类面孔情绪判断的一致性。我们系统操纵SE模块的压缩比、插入位置和预训练背景，并以40名参与者在悲伤—开心连续体上的判断作为比较基准。三个实验表明，在本研究所测试的配置中，面孔预训练、后序插入（FC7与FC8之间）且压缩比为16的SE-AlexNet与人类平均反应曲线的总体差异较小，但这一优势依赖预训练背景、压缩强度和插入位置的组合。Grad-CAM与面孔区域遮挡分析显示，该模型的分类输出对眼部等诊断性区域较敏感；这些结果反映模型显著性与输出敏感性，尚不能直接证明其具有与人类相同的注意机制。总体而言，表征压缩与特征重标定可能提高经典CNN与人类行为的对应程度，但结论仅适用于本研究训练的模型配置和所采样的人类观察者。未来需通过多随机种子训练、更广泛刺激及其他面孔判断任务检验其稳定性和可推广性。")

    replace_paragraph(doc.paragraphs[29],
        "图1  实验1的SE-AlexNet网络结构与主要结果。（a）SE-AlexNet架构及SE模块在AlexNet中的插入位置。除SE模块外，模型沿用AlexNet结构，输出层按本研究任务改为2个单元。（b）SE模块结构。FC7的4096维输出先按给定压缩比进行压缩，再恢复为4096维通道权重，并与FC7输出逐通道相乘后传入FC8。（c）模型输出的心理物理曲线：左为物体预训练模型，右为面孔预训练模型。（d）后序插入模型的PSE与人类平均PSE之差；同时给出AlexNet和VGG16基线。")

    replace_paragraph(doc.paragraphs[41],
        "为了将物体预训练与面孔预训练模型对齐到本研究的面孔情绪识别任务，我们采用迁移学习，将所有模型在同一情绪识别数据集上微调。该方法使模型能够复用预训练阶段获得的特征（Bengio, 2012）。")

    replace_paragraph(doc.paragraphs[57],
        "每个训练完成的神经网络在给定输入下产生确定性输出，因此每个模型配置只有一个固定PSE，不能形成模型内输出分布。为描述该固定PSE相对于人类群体分布的偏离程度，我们以40名参与者的PSE为参照进行单样本t检验。该检验回答的是“模型固定值偏离人类PSE分布的程度”，不应解释为模型本身具有被试间变异。")

    replace_paragraph(doc.paragraphs[63],
        "图1c分别给出面孔预训练与物体预训练模型的心理物理曲线，并以人类数据作为比较基准（黑色线段；mean PSE = 0.52, SEM = 0.01）。在所测试的后序插入SE-AlexNet中，压缩比为16的面孔预训练模型与人类平均反应曲线的RMSE较小。AlexNet与VGG16的反应曲线则在若干情绪强度上偏离人类平均曲线。")

    replace_paragraph(doc.paragraphs[69],
        "采用物体预训练权重的模型总体上与人类平均反应曲线差异较大。物体预训练的后序插入SE-AlexNet（压缩比为16；PSE = 0.73, t(39) = 10.02, p < 0.001, Cohen’s d = 1.58; RMSE = 0.24）及AlexNet（PSE = 0.67, t(39) = 6.83, p < 0.001, Cohen’s d = 1.08; RMSE = 0.27）的PSE均偏离人类基准。VGG16的PSE差异未达到0.05显著性水平（PSE = 0.49, t(39) = 1.89, p = 0.066, Cohen’s d = 0.30），但其反应曲线的RMSE仍较大（RMSE = 0.29）。这些结果说明，PSE差异检验与整条反应曲线的RMSE衡量的是不同方面，不宜互相替代。")

    replace_paragraph(doc.paragraphs[73],
        "大多数面孔预训练模型均具有较高的曲线拟合质量（见表2；例如压缩比为16的SE-AlexNet，R²fit = 0.99, RMSEfit = 0.04）。这说明所选心理物理函数能够较好描述模型输出的单调变化，但拟合质量本身不代表模型与人类行为一致。")

    replace_paragraph(doc.paragraphs[89],
        "其中，PSEi表示第i种神经网络配置产生的固定PSE，PSEhuman,j表示第j名参与者的PSE。ΔPSEij越小，表示该配置越接近相应参与者。该指标将每个固定模型值与同一组参与者逐一比较，形成可用于模型配置间比较的被试内误差指标；其推断范围仍限于本研究的参与者与固定模型输出。")

    replace_paragraph(doc.paragraphs[92],
        "结果表明，不同插入位置和压缩比的SE-AlexNet具有不同的PSE偏离模式。如图3a所示，前序插入模型均未同时达到接近人类平均PSE（0.52）和较低RMSE。AlexNet（PSE = 0.66, t(39) = 6.62, p < 0.001, Cohen’s d = 1.05; RMSE = 0.31）、VGG16（PSE = 0.84, t(39) = 7.55, p < 0.001, Cohen’s d = 1.19; RMSE = 0.16），以及压缩比为2（PSE = 0.64, t(39) = 5.78, p < 0.001, Cohen’s d = 0.91; RMSE = 0.15）、4（PSE = 0.92, t(39) = 16.91, p < 0.001, Cohen’s d = 2.67; RMSE = 0.32）、8（PSE = 0.94, t(39) = 16.99, p < 0.001, Cohen’s d = 2.69; RMSE = 0.33）、16（PSE = 0.37, t(39) = 5.21, p < 0.001, Cohen’s d = 0.82; RMSE = 0.17）和32（PSE = 0.28, t(39) = 7.82, p < 0.001, Cohen’s d = 1.24; RMSE = 0.19）的前序插入模型均不同程度偏离人类基准。高于或低于人类PSE代表不同方向的判断偏移，不能笼统解释为同一种正向偏差。")

    replace_paragraph(doc.paragraphs[103],
        "为解释三因素交互作用，我们在各插入位置与压缩比组合下比较两类预训练背景。差值定义为面孔预训练ΔPSE减去物体预训练ΔPSE，因此负值表示面孔预训练模型更接近人类，正值表示物体预训练模型更接近人类。在后序位置，8倍压缩比下面孔预训练模型的误差更高（Mdiff = 0.13, SE = 0.01, p < 0.001），而16倍（Mdiff = −0.08, SE = 0.01, p < 0.001）与32倍（Mdiff = −0.09, p < 0.001）压缩比下面孔预训练模型的误差更低。在前序位置，2倍压缩比下面孔预训练模型误差更低（Mdiff = −0.10, p < 0.001），4倍和8倍压缩比下物体预训练模型误差更低（ps < 0.001），16倍压缩比下差异不显著（Mdiff = −0.001, p = 0.393）。在中序位置，2倍（Mdiff = 0.06, p = 0.013）与8倍（Mdiff = 0.05, p < 0.001）压缩比下物体预训练模型误差更低；4倍压缩比下面孔预训练模型误差更低（Mdiff = −0.02, p < 0.001）；16倍压缩比下两类模型差异不显著（Mdiff = 0.003, p = 0.882）。")

    replace_paragraph(doc.paragraphs[104],
        "随后，我们在后序插入位置分别比较两类预训练模型的压缩比。面孔预训练模型在16倍压缩比下的ΔPSE低于2倍（Mdiff = −0.04, p = 0.003）和8倍（Mdiff = −0.14, p < 0.001），与4倍和32倍差异不显著（p = 0.348和p = 1.000）。物体预训练模型在8倍压缩比下的ΔPSE低于2倍（Mdiff = −0.09, p < 0.001）、4倍（Mdiff = −0.03, p < 0.001）、16倍（Mdiff = −0.07, p < 0.001）和32倍（Mdiff = −0.08, p < 0.001）。在后序插入、16倍压缩条件下，面孔预训练模型的ΔPSE也低于物体预训练模型（Mdiff = −0.08, p < 0.001）。这些比较表明，压缩比效应依赖预训练背景；“最优”仅指本研究所比较配置中的相对结果。")

    replace_paragraph(doc.paragraphs[105],
        "在16倍压缩比的面孔预训练模型中，Bonferroni校正后的两两比较未检出前序与后序位置差异（p = 1.000），前序与中序位置差异亦未达到0.05显著性水平（p = 0.060）。这些结果表示本研究未检出相应差异，但不能据此证明位置间等效或具有跨层级稳健性。物体预训练模型在16倍压缩比下，后序位置的ΔPSE高于前序位置（前序−后序：Mdiff = −0.08, p < 0.001），与中序位置的差异未达到0.05显著性水平（Mdiff = 0.10, p = 0.085）。在32倍压缩比下，前序与后序位置亦存在差异（p < 0.001）。")

    replace_paragraph(doc.paragraphs[128],
        "诊断比率取值为[0, 1]；数值越高，表示眼部与嘴部在三个ROI总显著性中所占比例越高。当三个ROI的PoS之和为0时，该比率未定义。全部模型、刺激与条件共6048条记录中有487条（8.05%）出现这一情况，计算诊断比率时予以剔除。因此，不同模型与条件的有效样本数可能不同，具体比较需报告相应n值。")

    replace_paragraph(doc.paragraphs[141],
        "在完整面孔条件下，后序插入、压缩比为16的面孔预训练SE-AlexNet具有较高的诊断比率（M = 0.95, SEM = 0.03），高于相同配置的物体预训练模型（M = 0.77, SEM = 0.03; Cohen’s d = 1.31, p < 0.001）。其眼部PoS也较高（M = 0.21 vs. 0.07, Cohen’s d = 0.94, p = 0.006），鼻部（M = 0.01 vs. 0.08, Cohen’s d = −1.32, p < 0.001）和嘴部（M = 0.03 vs. 0.12, Cohen’s d = −1.12, p < 0.001）PoS较低。诊断比率的差异主要反映较高的眼部权重与较低的鼻部权重；由于嘴部PoS同时下降，不能据此概括为模型对所有诊断性区域均增强。")

    replace_paragraph(doc.paragraphs[144],
        "图6  完整面孔条件下，眼部与嘴部显著性之和占眼、鼻、嘴三个ROI总显著性的比例（诊断比率）。（a）面孔预训练模型；（b）物体预训练模型。每个点为相应模型配置在21个情绪连续体刺激上的平均值，误差线为刺激间标准误；分母为0的记录不参与该比率计算。")

    replace_paragraph(doc.paragraphs[146],
        "我们进一步分析后序插入、面孔预训练、16倍压缩模型在遮挡条件下的PoS变化。以21张刺激为配对单位（df = 20），差值定义为ΔM = M完整 − M遮挡。眼部遮挡后，鼻部PoS的变化未达到0.05显著性水平（ΔM = −0.02, t(20) = −2.04, p = 0.055, Cohen’s d = −0.45），嘴部PoS无显著变化（ΔM = −0.002, t(20) = −0.14, p = 0.887, Cohen’s d = −0.03）。嘴部遮挡后，鼻部PoS升高（ΔM = −0.05, t(20) = −2.38, p = 0.028, Cohen’s d = −0.52），眼部PoS变化不显著（ΔM = −0.05, t(20) = −1.16, p = 0.259, Cohen’s d = −0.25）。鼻部遮挡后其余两个ROI的PoS均发生变化（ps < 0.040）。这些结果描述模型显著性分配对遮挡的敏感性，不直接等同于人类注意转移。")

    replace_paragraph(doc.paragraphs[151],
        "实验3显示，后序插入、压缩比为16的SE-AlexNet对眼部和嘴部遮挡较敏感，说明这些区域包含影响模型情绪分类输出的信息。Grad-CAM结果还显示该配置在完整面孔条件下具有较高的眼部显著性和诊断比率。上述结果与人类情绪识别中眼、嘴信息具有诊断价值的文献相容（Eisenbarth & Alpers, 2011; Iskra & Tomc, 2016），但模型显著性图不能被等同为人的注视或神经注意机制。")
    replace_paragraph(doc.paragraphs[152], "")

    replace_paragraph(doc.paragraphs[155],
        "本研究将SE模块嵌入AlexNet全连接层，系统操纵压缩比、插入位置与预训练背景，并从心理物理曲线、PSE偏差、遮挡敏感性和Grad-CAM显著性四个方面比较模型与人类行为。结果表明，压缩的作用并非普遍存在，而是依赖模型训练背景及其在网络中的位置。")

    replace_paragraph(doc.paragraphs[156],
        "在所测试的配置中，面孔预训练、后序插入且压缩比为16的SE-AlexNet与人类平均反应曲线的总体差异较小。三因素交互进一步表明，压缩比的作用随预训练背景和插入位置而变。Grad-CAM与遮挡分析则提供了分类输出依赖眼部等区域的证据。由于部分数值比较仍需与统一分析版本核对，且若干事后比较不显著，这些结果应被表述为特定配置下的行为对应，而非对“最优结构”或共同生物机制的证明。")
    replace_paragraph(doc.paragraphs[157], "")

    replace_paragraph(doc.paragraphs[158],
        "本研究的设计受到高级视觉皮层低维表征工作的启发（Chang et al., 2021; Chang & Tsao, 2017; Kriegeskorte et al., 2008; Young & Yamane, 1992）。在FC7的4096维输出上采用16倍压缩时，SE模块的瓶颈维度为256。该操作为检验“压缩与重标定是否影响模型—人类行为对应”提供了可控手段，但不能据此断言网络瓶颈与生物神经元数量或神经编码维度一一对应。")

    replace_paragraph(doc.paragraphs[159],
        "插入位置效应提示，特征重标定发生在网络的哪个阶段可能影响模型输出。后序插入位于较高层全连接表征之后，可在分类前重新加权高层特征；这一功能描述与高级视觉加工文献在概念上相容（Fusar-Poli et al., 2009; Vuilleumier & Pourtois, 2007），但现有实验没有直接建立网络层与脑区或加工阶段之间的同源对应。")

    replace_paragraph(doc.paragraphs[160],
        "Grad-CAM和遮挡实验补充了行为输出之外的证据：后序插入模型的分类显著性更多分配到眼部，遮挡眼部或嘴部也会改变模型反应曲线。然而，Grad-CAM反映的是输入区域对类别输出的局部敏感性，而不是眼动、主观注意或神经活动的直接测量。因此，本研究只能提出模型与人类诊断性特征使用之间的功能相似性，不能证明二者采用相同机制。")
    replace_paragraph(doc.paragraphs[161], "")

    replace_paragraph(doc.paragraphs[162],
        "这些结果也提示，判断模型是否“类人”需要区分不同层面：PSE衡量决策边界，RMSE衡量整条反应曲线差异，Grad-CAM和遮挡则衡量输出对空间区域的敏感性。单一指标不足以支持整体化加工、动态注意或生物合理性等更强结论；未来可将模型结果与眼动或神经表征直接比较。")

    replace_paragraph(doc.paragraphs[163],
        "SE模块在结构上较为轻量，因而可能具有工程应用价值。不过，本研究未系统报告参数量、FLOPs、推理延迟、能耗或多随机种子稳定性，尚不能据此得出计算效率或部署成本优于其他架构的结论。后续研究应在统一硬件与训练预算下同时比较行为对应、分类性能和资源消耗。")

    replace_paragraph(doc.paragraphs[164],
        "本研究的结论具有明确边界：任务仅涉及悲伤—开心判断，刺激的年龄、性别、种族等物理属性及可信度、吸引力、支配度等社会属性均未被系统操纵。未来可分别收集人类对这些属性的判断数据，建立各任务的行为基准，再检验不同压缩比和插入位置的模型与人类判断是否一致，以区分跨任务共性与任务特异性。还需在更广泛的CNN架构、刺激与自然情境中重复训练并比较不同随机种子，检验结果的稳定性和可推广性。")

    replace_paragraph(doc.paragraphs[167],
        "本研究表明，在AlexNet中加入SE模块后，模型与人类面孔情绪判断的对应程度会随压缩比、插入位置和预训练背景而变化。在本研究所测试的配置中，面孔预训练、后序插入且压缩比为16的模型总体上更接近人类平均反应曲线，并对眼部等区域表现出较高的输出敏感性。这些结果支持将表征压缩与特征重标定作为研究模型—人类行为对应的计算工具，但尚不足以证明模型复现了人类注意或神经压缩机制。")

    # Remove two redundant/poorly targeted references and their citations.
    for paragraph in list(doc.paragraphs):
        if paragraph.text.startswith("Selvaraju, R. R., Das, A., Vedantam"):
            delete_paragraph(paragraph)
        elif paragraph.text.startswith("Zhu, R., Qiu, T., Wang, J."):
            delete_paragraph(paragraph)
        elif "Selvaraju et al., 2016" in paragraph.text:
            replace_paragraph(paragraph, paragraph.text.replace("Selvaraju et al., 2016; ", ""))

    # Update English affiliations with city and postal code.
    english_units = {
        "1 Soochow University, Department of Psychology, China.":
            "1 Department of Psychology, Soochow University, Suzhou 215123, China.",
        "2 Beijing Normal University, State Key Laboratory of Cognitive Neuroscience and Learning, China.":
            "2 State Key Laboratory of Cognitive Neuroscience and Learning, Beijing Normal University, Beijing 100875, China.",
        "3 IDG/McGovern Institute, China.":
            "3 IDG/McGovern Institute for Brain Research, Beijing 100875, China.",
        "4 Peking University, School of Psychological and Cognitive Sciences, Beijing Key Laboratory of Behavior and Mental Health, China.":
            "4 School of Psychological and Cognitive Sciences, Beijing Key Laboratory of Behavior and Mental Health, Key Laboratory of Machine Perception (Ministry of Education), Peking University, Beijing 100871, China.",
        "5 Peking University, Institute for Artificial Intelligence, China.":
            "5 Institute for Artificial Intelligence, Peking University, Beijing 100871, China.",
        "6 Beijing Institute for General Artificial Intelligence, State Key Laboratory of General Artificial Intelligence, China.":
            "6 Beijing Institute for General Artificial Intelligence, State Key Laboratory of General Artificial Intelligence, Beijing 100080, China.",
    }
    for paragraph in doc.paragraphs:
        if paragraph.text in english_units:
            replace_paragraph(paragraph, english_units[paragraph.text])

    # Journal typography.
    for i, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if i == 0:
            style_paragraph(paragraph, "黑体", 22, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        elif i == 1:
            style_paragraph(paragraph, "仿宋", 14, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        elif 2 <= i <= 7:
            style_paragraph(paragraph, "宋体", 9, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        elif text == "摘要" or text == "参考文献":
            style_paragraph(paragraph, "黑体", 10.5, bold=True)
        elif text.startswith("图"):
            style_paragraph(paragraph, "宋体", 7.5, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        elif text.startswith("表"):
            style_paragraph(paragraph, "宋体", 9, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        elif text and text[0].isdigit() and any(text.startswith(f"{n} ") for n in range(1, 9)):
            style_paragraph(paragraph, "宋体", 14, bold=False)
        elif text and text[0].isdigit() and "." in text.split(" ", 1)[0]:
            style_paragraph(paragraph, "黑体", 10.5, bold=True)
        else:
            style_paragraph(paragraph, "宋体", 10.5)

    # English back matter has its own sizes.
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "A squeeze-and-excitation computational model for efficient face cognition":
            style_paragraph(paragraph, "Times New Roman", 14, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        elif text.startswith("BAO Jiayu"):
            style_paragraph(paragraph, "Times New Roman", 10.5, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        elif text.startswith(tuple(f"{i} " for i in range(1, 7))) and "China." in text:
            style_paragraph(paragraph, "Times New Roman", 7.5, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        elif text in {"Abstract"} or text.startswith("Keywords "):
            style_paragraph(paragraph, "Times New Roman", 10.5, bold=(text == "Abstract"))

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    style_paragraph(paragraph, "宋体", 9, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    # Author-facing comments for issues that cannot be solved by wording alone.
    add_comment(doc, "batch size为64，学习率为",
                "两处学习率均缺失。请从最终训练配置或日志填入精确数值，并分别核对身份预训练与情绪微调阶段；不能只保留Krizhevsky等人的引文。")
    add_comment(doc, "本研究使用Wichmann和Hill",
                "稿件公式是四参数logistic，但当前公开拟合代码使用基于erf的累积高斯函数，README也存在版本差异。请确定录用稿实际采用的函数，并同步公式、方法、代码和结果。")
    add_comment(doc, "mean PSE = 0.52, SEM = 0.01",
                "现有40人PSE文件重算为M=0.5143、SD=0.0924、SEM=0.0146；稿件表1的若干t和d也无法由该分布与表中PSE复现。请锁定唯一分析版本后统一替换正文、表1、图1、图3、表3与英文摘要。")
    add_comment(doc, "PSE = 0.61, t(39) = 1.33",
                "t(39)=1.33在普通双侧检验下p约为0.192，且单样本Cohen’s d通常约为0.21；原稿p=0.220、d=0.47不能同时由同一检验得到。须回到原始脚本重算，不能用倒推值直接替换。")
    add_comment(doc, "F(4, 156) = 8.81, p = 0.003",
                "未校正F(4,156)=8.81的p远小于0.003；如p来自Greenhouse–Geisser等校正，应报告校正后的自由度/epsilon并在方法中说明。当前归档结果没有对应的简单交互输出。")
    add_comment(doc, "M = 0.95, SEM = 0.03",
                "这两组诊断比率可由GradCAM_Results.csv的完整面孔条件复现（各n=21）。但Cohen’s d与大部分p值更接近把同一批ImageID当作独立样本的检验；请确认是否应采用配对检验，并说明多重比较校正。")
    add_comment(doc, "4096维通道权重",
                "正文已按SE逐通道乘法逻辑改为4096维，但图1b图片仍标注9216。请从源图同步改为4096；16倍瓶颈应为4096/16=256，而不是约200。")
    add_comment(doc, "误差线为刺激间标准误",
                "图6使用21个情绪连续体刺激，而行为实验与PSE拟合使用9个刺激。请在方法中明确21点连续体的生成、其与9点行为刺激的关系及仅用于Grad-CAM/ROI分析的原因。")

    doc.save(MAIN_OUT)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(text)
    size = 14 if level == 1 else 11
    style_paragraph(p, "Hiragino Sans GB", size, bold=True)
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(4)


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, "Hiragino Sans GB", 10.5, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, "Hiragino Sans GB", 10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, "Hiragino Sans GB", 10.5)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)


def build_response_working_draft() -> None:
    doc = Document(RESPONSE_SOURCE)
    configure_page(doc)
    doc.add_page_break()
    add_heading(doc, "录用后编辑意见逐条回应（工作稿，待作者核定）", 1)
    add_body(doc, "说明：以下回应依据2026年10月14日对录用稿、分析仓库和Grad-CAM结果表的核查形成。涉及统计源版本、学习率、心理物理函数及英文润色证明的项目尚未闭环，因此本稿不能直接随终稿发送。")

    responses = [
        ("1. 精减稿件", "已重写中文摘要，并压缩总讨论与结论中的重复表述；同时将“复现人类机制”等强结论降为模型—行为对应层面的表述。待统计版本锁定后还需进行一次终读。"),
        ("2. 关键词格式", "已删除中英文“关键词/Keywords”后的冒号；中文关键词统一用逗号，英文关键词统一用英文逗号分隔。"),
        ("3. 伦理委员会", "已恢复为“苏州大学伦理委员会（SUDA20210402H03）”。"),
        ("4. 统计精度", "已统一常规统计量保留两位、p值保留三位并补足前导0。但核查发现若干统计值与现有40人数据/脚本不属于同一分析版本，必须重算并全稿同步后才能确认完成。"),
        ("5. 小数点前0", "已核查正文与表格的统计表达，补足小于1的统计量前导0。"),
        ("6. 图表题格式", "已删除图表编号后的点号并将图表标题居中；表格字体按小五号、插图说明按六号设置。"),
        ("7. 黑白印刷区分度", "现有图中仍有多条近似色曲线。建议在源图中同时使用线型、点型和直接标签，并打印灰度稿复核；此项尚需源图修改。"),
        ("8. 全文语言核查", "已修正“真是人类知觉”“FC7FC8”“t=11.11. p”等错误，并重写编辑指出的未来研究句。对Grad-CAM、人类注意、生物机制与计算效率的过强推断也已收敛。"),
        ("9. 正文引文", "程序核对显示原稿79条文献均能在正文找到对应引用；修订中删除了重复的Grad-CAM预印本与一条过于偏题的迁移学习文献，并同步删除正文引文。最终统计修订后需再跑一次双向核对。"),
        ("10. 文后参考文献", "已核对题名句式大小写、期刊名大小写和页码范围连接号；会议文献保留会议信息。无DOI或数据库型条目仍需作者依据原始来源人工确认。"),
        ("11. 英文润色", "英文标题与摘要已按学术英语重写并与中文结论边界保持一致，但本次内部修改不能替代期刊要求的母语学者信息或专业机构润色证明。请作者补充相应证明材料。"),
        ("12. 全面一致性核查", "已核查样本量、刺激数、PSE/SEM、t/p/d、ANOVA、Grad-CAM分母和代码版本。发现统计版本不一致、9/21个刺激说明不足、心理物理函数与代码不一致等问题，详见《最终核查报告》；在这些问题闭环前不能写“已全部核对无误”。"),
        ("13. 中英文姓名与单位", "中英文姓名与单位均置于相应题目下，英文姓氏大写；英文单位补充了城市和邮编。"),
        ("14. 首页脚注", "首页已包含投稿日期2026-04-09、基金号32200850和通信作者E-mail。请通信作者最终确认基金号、投稿日期和邮箱是否为期刊登记版本。"),
        ("15. 终稿格式", "已按通知设置题名、作者、单位、摘要、正文、标题、图表、参考文献及英文部分字号，数字和英文使用Times New Roman，全稿1.5倍行距。由于本机缺少期刊指定的全部中文字体，仍须在装有相应字体的Word中打印复核分页与公式。"),
    ]
    for title, response in responses:
        add_heading(doc, title, 2)
        add_body(doc, f"回应：{response}", bold_prefix="回应：")

    add_heading(doc, "ChinaXiv提交状态（作者二选一后保留其一）", 2)
    add_body(doc, "□ 已在投稿时提交，现已用最新录用稿替换原稿。")
    add_body(doc, "□ 投稿时未提交，现已将录用稿提交至中国心理学预印本平台（http://psych.chinaxiv.org）。")

    for paragraph in doc.paragraphs:
        if paragraph.text and "录用后编辑意见逐条回应" not in paragraph.text:
            for run in paragraph.runs:
                if run.font.name is None:
                    set_run_font(run, "Hiragino Sans GB", 10.5)
    doc.save(RESPONSE_OUT)


def add_bullet(doc: Document, text: str, *, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    r = p.add_run(text)
    set_run_font(r, "Hiragino Sans GB", 10.5)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(2)


def build_audit_report() -> None:
    doc = Document()
    configure_page(doc)
    styles = doc.styles
    if "List Bullet 2" not in [s.name for s in styles]:
        styles.add_style("List Bullet 2", WD_STYLE_TYPE.PARAGRAPH)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("xb26-0201 录用终稿逐条核查报告")
    set_run_font(r, "Hiragino Sans GB", 18, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("核查日期：2026-10-14｜结论：当前仍不宜直接返回期刊")
    set_run_font(r, "Hiragino Sans GB", 11, bold=True, color=RGBColor(192, 0, 0))

    add_heading(doc, "一、先给结论", 1)
    add_body(doc, "文字、格式、关键词、伦理信息、作者单位和大部分引文格式可以直接修正；Grad-CAM完整面孔条件下的诊断比率也能从原始结果表复现。但核心PSE统计来自与现有仓库不同的分析版本，且方法中仍缺两个学习率，心理物理函数与公开代码不一致。因此，本次生成的是“作者核定版”，不是可直接发送的干净终稿。")

    add_heading(doc, "二、必须在回稿前闭环的红色问题", 1)
    blockers = [
        "锁定唯一PSE分析版本。现有40人PSE文件重算为M=0.5143、SD=0.0924、SEM=0.0146；稿件表1中的VGG16/AlexNet等t与d无法由这组分布和表中PSE复现。仓库中的多个Fitting_Results/Analysis_Results文件也与稿件数值不一致。",
        "重算并同步所有PSE推断统计。最明显的逻辑冲突是PSE=0.61、t(39)=1.33、p=0.220、d=0.47：常规双侧t检验下p约为0.192，d约为0.21。不能只改一项。",
        "确定心理物理函数。稿件公式为四参数logistic，当前公开代码使用基于erf的累积高斯函数；9点/21点刺激和不同输出文件也显示存在版本分叉。",
        "补齐两个训练阶段的学习率，并用最终训练配置核对batch size、epoch、冻结层、随机种子和权重文件。",
        "核定ANOVA简单交互。F(4,156)=8.81对应的未校正p远小于0.003；若使用球形性校正，需要报告校正后的自由度或epsilon。",
        "核定ROI比较的检验设计。正文d值可由每组21张图的独立样本效应量复现，但同一ImageID跨模型天然配对；应由作者确认配对/独立选择并处理多重比较。",
        "修改图1源图：SE激励后维度应恢复为4096以便逐通道相乘；4096/16=256，不是图/讨论中的9216或约200。",
    ]
    for item in blockers:
        add_bullet(doc, item)

    add_heading(doc, "三、已经独立复核并可保留的结果", 1)
    verified = [
        "GradCAM_Results.csv共6048条，三个ROI全零487条，占8.05%，与稿件一致。",
        "后序插入、面孔预训练、R16、完整面孔条件的诊断比率：n=21，M=0.9495，SEM=0.0310；对应物体预训练模型：n=21，M=0.7736，SEM=0.0276。正文四舍五入为0.95±0.03和0.77±0.03，可复现。",
        "ANOVA归档的三因素交互F(8,312)=26.92，偏η²约0.41，可复现；但简单交互的校正细节仍缺。",
        "PostHoc_3Way_Compare_Pretrain.csv明确将差值定义为Face−Object。原稿中序位置R2/R8/R4的优劣方向写反，作者核定版已按差值方向改正。",
        "原稿79条参考文献均能在正文找到对应第一作者与年份；修订后删除2条冗余/偏题文献，并同步删除引文。",
    ]
    for item in verified:
        add_bullet(doc, item)

    add_heading(doc, "四、15项编辑要求落实表", 1)
    rows = [
        ("1", "精减", "部分完成", "摘要、总讨论和结论已明显压缩；统计定稿后再终读。"),
        ("2", "关键词", "已完成", "无冒号；中英文均用逗号。"),
        ("3", "伦理", "已完成", "苏州大学伦理委员会（SUDA20210402H03）。"),
        ("4–5", "统计精度/前导0", "格式完成，数值未闭环", "显示格式已统一；推断统计须重算。"),
        ("6", "图表题", "已完成", "无点号、居中。"),
        ("7", "黑白区分", "待源图修改", "建议线型+点型+直接标签并打印灰度稿。"),
        ("8", "通读订误", "主要文字已改", "包括编辑点名句、错字、方向性和过强结论。"),
        ("9–10", "引用与文献", "内部对应完成", "无孤立条目；无DOI来源仍需人工核原文。"),
        ("11", "英文润色证明", "作者事项", "内部润色不能替代母语学者信息/公司证明。"),
        ("12", "全面一致性", "发现阻断问题", "详见第二部分。"),
        ("13", "姓名与单位", "已完成", "英文姓氏大写，英文单位补城市邮编。"),
        ("14", "首页脚注", "已存在，待确认", "日期2026-04-09、基金32200850、邮箱hjying@suda.edu.cn。"),
        ("15", "版式", "已程序设置", "仍需在Word中打印核页、公式和中文字体。"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, label in enumerate(["序号", "要求", "状态", "说明"]):
        hdr[i].text = label
        set_cell_shading(hdr[i], "D9EAF7")
    set_repeat_table_header(table.rows[0])
    for no, requirement, status, note in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, [no, requirement, status, note]):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if "完成" in status and "未" not in status and "部分" not in status:
            set_cell_shading(cells[2], "E2F0D9")
        elif "作者事项" in status or "待" in status or "阻断" in status or "未闭环" in status:
            set_cell_shading(cells[2], "FCE4D6")
        else:
            set_cell_shading(cells[2], "FFF2CC")
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                style_paragraph(paragraph, "Hiragino Sans GB", 9)

    add_heading(doc, "五、建议的最短闭环顺序", 1)
    steps = [
        "由负责统计的作者指定“唯一真源”：最终40人PSE文件、模型PSE输出、拟合函数、ANOVA输出及生成脚本。",
        "从该真源一次性导出表1、表2、表3和图1/3/4/6所需数值，禁止在Word中手工拼接不同版本。",
        "重新计算并程序化导出M/SD/SEM、t、df、p、Cohen’s d、RMSE、F、校正df/epsilon、偏η²及事后比较。",
        "同步更新中英文摘要、正文、图表、审稿回应和公开仓库；之后再跑引用双向核对与数值全文搜索。",
        "在Word中以黑白打印：重点检查图1维度、曲线线型、图6说明、表格分页、公式、脚注和中英文字体。",
        "补英文润色证明并确认ChinaXiv状态后，再生成无批注、无高亮的提交版。",
    ]
    for idx, step in enumerate(steps, 1):
        p = doc.add_paragraph()
        r1 = p.add_run(f"{idx}. ")
        set_run_font(r1, "Hiragino Sans GB", 10.5, bold=True)
        r2 = p.add_run(step)
        set_run_font(r2, "Hiragino Sans GB", 10.5)
        p.paragraph_format.line_spacing = 1.4
        p.paragraph_format.space_after = Pt(3)

    add_heading(doc, "六、交付文件说明", 1)
    add_body(doc, "《正文_作者核定版_勿直接提交》：已完成安全的文字、方向性、格式和单位修订，并在关键阻断位置加入Word批注。")
    add_body(doc, "《审稿意见及说明_录用后补充工作稿》：在原三轮回应后补入本次15项录用后编辑意见的逐条回应草稿。")
    add_body(doc, "本报告：用于作者内部逐项关账。所有原始文件均未覆盖。")

    doc.save(REPORT_OUT)


if __name__ == "__main__":
    edit_main_manuscript()
    build_response_working_draft()
    build_audit_report()
    for path in (MAIN_OUT, RESPONSE_OUT, REPORT_OUT):
        print(path)
